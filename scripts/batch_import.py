import json, requests, time
from openpyxl import load_workbook

# 1. 读取 token
with open('C:/Users/Nick/Desktop/teamwork-clone/token.json', 'r', encoding='utf-8') as f:
    token_data = json.load(f)
token = token_data.get('access_token', '')

if not token:
    print("❌ Token 为空，请先登录")
    exit(1)

print(f"✓ Token 已读取")

# 2. 读取 Excel
wb = load_workbook('C:/Users/Nick/Desktop/Teamwork_待匯入地址_完整欄位.xlsx')
ws = wb.active

# 找到所有未导入的数据（AL列为空且有序号）
rows_to_import = []
for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if row[0] and isinstance(row[0], int):
        # 检查 AL 列（第38列，索引37）是否已标记
        al_value = row[37] if len(row) > 37 else None
        if al_value != '√':
            rows_to_import.append((row_idx, row))

print(f"✓ 找到 {len(rows_to_import)} 笔待导入数据")

if not rows_to_import:
    print("✓ 所有数据已导入完成！")
    exit(0)

# 3. API 配置
API_BASE = 'https://hk-teamwork-api-rp.transnational-grp.com'
headers = {
    'apikey': '5567GGH67225HYVGG',
    'Authorization': f'Bearer {token}',
    'Origin': 'https://hk-teamwork.transnational-grp.com/',
    'Referer': 'https://hk-teamwork.transnational-grp.com/',
    'Content-Type': 'application/json'
}

# 4. 分批导入（每批100笔，间隔60秒）
BATCH_SIZE = 100
BATCH_INTERVAL = 60  # 秒

total_success = 0
total_fail = 0
batch_num = 0

for batch_start in range(0, len(rows_to_import), BATCH_SIZE):
    batch_num += 1
    batch_end = min(batch_start + BATCH_SIZE, len(rows_to_import))
    batch_rows = rows_to_import[batch_start:batch_end]
    
    print(f"\n{'='*60}")
    print(f"批次 {batch_num}: 导入第 {batch_start+1}-{batch_end} 笔（共 {len(batch_rows)} 笔）")
    print(f"{'='*60}")
    
    batch_success = 0
    batch_fail = 0
    
    for i, (row_idx, row) in enumerate(batch_rows, 1):
        seq = row[0]
        
        # 构建 payload
        payload = {
            'AddressName': row[2] or '',
            'CostCenterName': row[3] or '',
            'ContactName': row[4] or '-',
            'Countryid': str(row[9]) if row[9] else '1',
            'CityId': str(row[10]) if row[10] else '',
            'Districtid': str(row[11]) if row[11] else '',
            'Zoneid': str(row[12]) if row[12] else '',
            'Stateid': str(row[13]) if row[13] else '2',
            'Regionid': row[14] or '',
            'PostalCode': row[15] or '',
            'Address1': row[5] or '',
            'Address2': row[6] or '',
            'Address3': row[7] or '',
            'Address4': row[8] or '',
            'Phone1': row[16] or '-',
            'Ext1': row[17] or '',
            'Phone2': row[18] or '',
            'Ext2': row[19] or '',
            'Mobile': row[20] or '',
            'Email': row[21] or '-',
            'ALatitude': row[22] or '',
            'ALongitude': row[23] or '',
            'Collection': 'True',
            'Delivery': 'True',
            'Collection_Instruction': '',
            'DefaultAddress': 'false',
            'DefaultShipping': 'no',
            'DefaultBilling': 'no',
            'BranchName': '',
            'Department': '',
            'CBranchId': '',
            'Contactid': 4323,
            'CustomerId': 148551,
            'CD': 'HKTeamwork'
        }
        
        try:
            resp = requests.post(
                f'{API_BASE}/api/Contact/SavecontactAddress',
                headers=headers,
                json=payload,
                timeout=30
            )
            result = resp.json()
            
            if result.get('IsSuccess'):
                batch_success += 1
                total_success += 1
                # 在 AL 列（第38列）写入 √
                ws.cell(row=row_idx, column=38, value='√')
                if i % 20 == 0:
                    print(f"  [{i}/{len(batch_rows)}] ✓ 本批已导入 {batch_success} 笔")
            else:
                batch_fail += 1
                total_fail += 1
                print(f"  [{i}] ✗ 序号 {seq} 失败: {result.get('Message')}")
        
        except Exception as e:
            batch_fail += 1
            total_fail += 1
            print(f"  [{i}] ✗ 序号 {seq} 异常: {e}")
        
        # 每 20 笔暂停 1 秒，避免过快
        if i % 20 == 0:
            time.sleep(1)
    
    print(f"\n批次 {batch_num} 完成:")
    print(f"  ✓ 成功: {batch_success} 笔")
    print(f"  ✗ 失败: {batch_fail} 笔")
    
    # 保存 Excel（每批次后保存）
    wb.save('C:/Users/Nick/Desktop/Teamwork_待匯入地址_完整欄位.xlsx')
    print(f"  ✓ Excel 已保存")
    
    # 如果还有下一批，等待60秒
    if batch_end < len(rows_to_import):
        print(f"\n⏳ 等待 60 秒后继续下一批...")
        time.sleep(BATCH_INTERVAL)

print(f"\n{'='*60}")
print(f"全部导入完成!")
print(f"{'='*60}")
print(f"总成功: {total_success} 笔")
print(f"总失败: {total_fail} 笔")
print(f"总批次: {batch_num} 批")
